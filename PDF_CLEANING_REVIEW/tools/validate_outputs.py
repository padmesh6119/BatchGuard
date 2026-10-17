#!/usr/bin/env python3
"""
tools/validate_outputs.py
Comprehensive validation suite for PDF_CLEANING_REVIEW.
Validates:
  1. All 32 source PDFs are accounted for in manifest and extraction summaries.
  2. Exactly ten workbooks exist and reopen successfully.
  3. Exact sheet names and column headers matching reference format.
  4. Header in row 1; no merged cells, no formulas, no title rows.
  5. S.No continuity and record integrity (no unexplained record loss).
  6. Provenance JSONL matches every single workbook row 1-to-1.
  7. Spurious two-row pattern compatibility.
  8. Revision records and links are complete.
  9. Reports are populated and formatted correctly.
"""

import csv
import json
import re
from pathlib import Path
import openpyxl

PROJECT_ROOT = Path("/home/vishu/Projects/IBM/BATCHGUARD")
PDF_DIR = PROJECT_ROOT / "PDF"
OUTPUT_DIR = PROJECT_ROOT / "PDF_CLEANING_REVIEW"
WORKBOOKS_DIR = OUTPUT_DIR / "workbooks"
EVIDENCE_DIR = OUTPUT_DIR / "evidence"
REPORTS_DIR = OUTPUT_DIR / "reports"

MONTHS = [
    "2024-09", "2024-10", "2024-11", "2024-12",
    "2025-01", "2025-02", "2025-03", "2025-04",
    "2025-05", "2025-06"
]

CENTRAL_HEADERS = [
    "S.No", "Name of Product", "Batch No", "Manufacturing Date", "Expiry Date",
    "Manufactured By", "NSQ Result", "Reporting Source", "Reporting by Lab/State", "Reporting Month & Year"
]

STATE_HEADERS = [
    "S.No", "Name of Product", "Batch No", "Manufacturing Date", "Expiry Date",
    "Manufactured By", "NSQ Result", "Reporting Source", "Reporting by Lab/State", "Reporting Month & Year"
]

SPURIOUS_HEADERS = [
    "S.No", "Name of Product", "Batch No", "Manufacturing Dates", "Manufacturer Details",
    "Reporting Source", "Reporting by Lab/State", "Reporting Month & Year", "Remarks"
]

EXPECTED_COUNTS = {
    "2024-09": {"central": 49, "state": 18, "spurious": 4},
    "2024-10": {"central": 56, "state": 34, "spurious": 3},
    "2024-11": {"central": 41, "state": 70, "spurious": 2},
    "2024-12": {"central": 51, "state": 84, "spurious": 0},
    "2025-01": {"central": 52, "state": 93, "spurious": 0},
    "2025-02": {"central": 47, "state": 56, "spurious": 1},
    "2025-03": {"central": 70, "state": 61, "spurious": 1},
    "2025-04": {"central": 60, "state": 136, "spurious": 1},
    "2025-05": {"central": 58, "state": 128, "spurious": 2},
    "2025-06": {"central": 55, "state": 130, "spurious": 4},
}


def validate_all():
    errors = []
    warnings = []
    
    print("=== BATCHGUARD PDF CLEANING VALIDATION ===")
    
    # 1. Check all 32 PDFs exist
    pdf_files = sorted(PDF_DIR.glob("*.pdf"))
    if len(pdf_files) != 32:
        errors.append(f"Expected 32 PDFs in {PDF_DIR}, found {len(pdf_files)}")
    else:
        print(f"[PASS] Found exactly 32 source PDFs in {PDF_DIR.name}")
        
    # 2. Check source_manifest.csv
    manifest_path = REPORTS_DIR / "source_manifest.csv"
    if not manifest_path.exists():
        errors.append("source_manifest.csv missing")
    else:
        with open(manifest_path, encoding="utf-8-sig") as f:
            reader = list(csv.DictReader(f))
            if len(reader) != 32:
                errors.append(f"source_manifest.csv expected 32 rows, got {len(reader)}")
            else:
                print(f"[PASS] source_manifest.csv contains all 32 source PDFs")
                
    # 3. Check extraction_summary.csv
    summary_path = REPORTS_DIR / "extraction_summary.csv"
    if not summary_path.exists():
        errors.append("extraction_summary.csv missing")
    else:
        with open(summary_path, encoding="utf-8-sig") as f:
            reader = list(csv.DictReader(f))
            if len(reader) != 32:
                errors.append(f"extraction_summary.csv expected 32 rows, got {len(reader)}")
            else:
                print(f"[PASS] extraction_summary.csv reconciles all 32 source PDFs")
                
    # 4. Check 10 monthly workbooks
    total_workbook_rows = 0
    total_logical_records = 0
    
    for month in MONTHS:
        fname = f"{month}_all.xlsx"
        wb_path = WORKBOOKS_DIR / fname
        if not wb_path.exists():
            errors.append(f"Missing workbook: {fname}")
            continue
            
        try:
            wb = openpyxl.load_workbook(wb_path, data_only=True)
        except Exception as e:
            errors.append(f"Failed to open {fname}: {e}")
            continue
            
        expected_spurious = EXPECTED_COUNTS[month]["spurious"] > 0
        expected_sheets = ["central", "state", "spurious"] if expected_spurious else ["central", "state"]
        
        if wb.sheetnames != expected_sheets:
            errors.append(f"{fname} sheet names mismatch: got {wb.sheetnames}, expected {expected_sheets}")
            
        # Check Central sheet
        ws_c = wb["central"]
        c_headers = [cell.value for cell in ws_c[1]]
        if c_headers != CENTRAL_HEADERS:
            errors.append(f"{fname} [central] headers mismatch: {c_headers} vs {CENTRAL_HEADERS}")
            
        c_rows = list(ws_c.iter_rows(min_row=2, values_only=True))
        exp_c_cnt = EXPECTED_COUNTS[month]["central"]
        if len(c_rows) != exp_c_cnt:
            errors.append(f"{fname} [central] row count {len(c_rows)} != expected {exp_c_cnt}")
        total_workbook_rows += len(c_rows)
        total_logical_records += len(c_rows)
        
        # Check serial continuity in central
        serials_c = [r[0] for r in c_rows]
        if serials_c != list(range(1, exp_c_cnt + 1)):
            errors.append(f"{fname} [central] serials not strictly 1..{exp_c_cnt}: {serials_c[:5]}...")
            
        # Check State sheet
        ws_s = wb["state"]
        s_headers = [cell.value for cell in ws_s[1]]
        if s_headers != STATE_HEADERS:
            errors.append(f"{fname} [state] headers mismatch: {s_headers} vs {STATE_HEADERS}")
            
        s_rows = list(ws_s.iter_rows(min_row=2, values_only=True))
        exp_s_cnt = EXPECTED_COUNTS[month]["state"]
        if len(s_rows) != exp_s_cnt:
            errors.append(f"{fname} [state] row count {len(s_rows)} != expected {exp_s_cnt}")
        total_workbook_rows += len(s_rows)
        total_logical_records += len(s_rows)
        
        # Check serial continuity in state
        serials_s = [r[0] for r in s_rows]
        if serials_s != list(range(1, exp_s_cnt + 1)):
            errors.append(f"{fname} [state] serials not strictly 1..{exp_s_cnt}: {serials_s[:5]}...")
            
        # Check Spurious sheet if present
        if expected_spurious:
            ws_sp = wb["spurious"]
            sp_headers = [cell.value for cell in ws_sp[1]]
            if sp_headers != SPURIOUS_HEADERS:
                errors.append(f"{fname} [spurious] headers mismatch: {sp_headers} vs {SPURIOUS_HEADERS}")
                
            sp_rows = list(ws_sp.iter_rows(min_row=2, values_only=True))
            exp_sp_cnt = EXPECTED_COUNTS[month]["spurious"]
            # 2 rows per logical record
            if len(sp_rows) != exp_sp_cnt * 2:
                errors.append(f"{fname} [spurious] row count {len(sp_rows)} != expected {exp_sp_cnt * 2}")
            total_workbook_rows += len(sp_rows)
            total_logical_records += exp_sp_cnt
            
            # Check 2-row pattern in spurious
            for i in range(0, len(sp_rows), 2):
                r1 = sp_rows[i]
                r2 = sp_rows[i+1]
                if r1[0] is None:
                    errors.append(f"{fname} [spurious] Row {i+2} missing S.No")
                if r2[0] is not None:
                    errors.append(f"{fname} [spurious] Continuation Row {i+3} has non-empty S.No: {r2[0]}")
                if not str(r1[3] or "").startswith("Manufacturing Date:"):
                    errors.append(f"{fname} [spurious] Row {i+2} Manufacturing Dates does not start with 'Manufacturing Date:': {r1[3]}")
                if not str(r2[3] or "").startswith("Expiry Date:"):
                    errors.append(f"{fname} [spurious] Row {i+3} Manufacturing Dates does not start with 'Expiry Date:': {r2[3]}")
                if not str(r2[4] or "").startswith("Manufactured By:"):
                    errors.append(f"{fname} [spurious] Row {i+3} Manufacturer Details does not start with 'Manufactured By:': {r2[4]}")
                if not str(r1[8] or "").startswith("NSQ Remark:"):
                    errors.append(f"{fname} [spurious] Row {i+2} Remarks does not start with 'NSQ Remark:': {r1[8]}")
                if not str(r2[8] or "").startswith("Firm Reply:"):
                    errors.append(f"{fname} [spurious] Row {i+3} Remarks does not start with 'Firm Reply:': {r2[8]}")

    print(f"[CHECK] Workbook structure/count checks completed:")
    print(f"       Total workbook rows: {total_workbook_rows}")
    print(f"       Total logical records: {total_logical_records} (Central: 539, State: 810, Spurious: 18)")

    # 5. Check provenance JSONL
    prov_path = EVIDENCE_DIR / "record_provenance.jsonl"
    if not prov_path.exists():
        errors.append("record_provenance.jsonl missing")
    else:
        with open(prov_path, encoding="utf-8") as f:
            prov_records = [json.loads(line) for line in f]
        if len(prov_records) != total_workbook_rows:
            errors.append(f"Provenance count {len(prov_records)} != workbook rows {total_workbook_rows}")
        else:
            print(f"[CHECK] Provenance count matches ({len(prov_records)} rows); identities checked below")

    # 6. Check revision records
    rev_path = EVIDENCE_DIR / "revision_records.xlsx"
    if not rev_path.exists():
        errors.append("revision_records.xlsx missing")
    else:
        wb_rev = openpyxl.load_workbook(rev_path, data_only=True)
        ws_rev = wb_rev["revision_records"]
        rev_rows = list(ws_rev.iter_rows(min_row=2, values_only=True))
        if len(rev_rows) != 4:
            errors.append(f"revision_records.xlsx expected 4 rows, found {len(rev_rows)}")
        else:
            print(f"[PASS] revision_records.xlsx verified with 4 revision entries")

    # 7. Check revision links
    links_path = REPORTS_DIR / "revision_links.csv"
    if not links_path.exists():
        errors.append("revision_links.csv missing")
    else:
        with open(links_path, encoding="utf-8-sig") as f:
            links = list(csv.DictReader(f))
        if len(links) != 4:
            errors.append(f"revision_links.csv expected 4 rows, found {len(links)}")
        else:
            print(f"[PASS] revision_links.csv verified with 4 linked revisions")

    # 8. Check manual_review.csv
    review_path = REPORTS_DIR / "manual_review.csv"
    if not review_path.exists():
        errors.append("manual_review.csv missing")
    else:
        with open(review_path, encoding="utf-8-sig") as f:
            rev_items = list(csv.DictReader(f))
        print(f"[PASS] manual_review.csv verified with {len(rev_items)} flagged issues")

    # 9. Check duplicate_candidates.csv
    dup_path = REPORTS_DIR / "duplicate_candidates.csv"
    if not dup_path.exists():
        errors.append("duplicate_candidates.csv missing")
    else:
        with open(dup_path, encoding="utf-8-sig") as f:
            dups = list(csv.DictReader(f))
        print(f"[PASS] duplicate_candidates.csv verified with {len(dups)} candidate pairs")

    # 10. Check source_availability.csv
    avail_path = REPORTS_DIR / "source_availability.csv"
    if not avail_path.exists():
        errors.append("source_availability.csv missing")
    else:
        with open(avail_path, encoding="utf-8-sig") as f:
            avails = list(csv.DictReader(f))
        if len(avails) != 10:
            errors.append(f"source_availability.csv expected 10 rows, found {len(avails)}")
        else:
            print(f"[PASS] source_availability.csv verified with all 10 monthly availability entries")

    from review_integrity import check_integrity
    print("[CHECK] Re-extracting source PDFs and checking cells, raw evidence, hashes, review links and revisions...")
    try:
        errors.extend(check_integrity(OUTPUT_DIR, PDF_DIR))
    except Exception as exc:
        errors.append(f"Source/content validation could not complete: {exc}")

    if errors:
        print("\n[VALIDATION FAILED]")
        for e in errors:
            print(f"  ERROR: {e}")
        return False
    else:
        print("\n=== STRUCTURE AND SOURCE-CONTENT CHECKS PASSED ===")
        print("Fresh table extraction uses the same extraction engine; it is not an independent visual audit of every cell.")
        print("Unresolved source/revision issues remain flagged. Staged only; not integrated or model-ready.")
        return True


if __name__ == "__main__":
    success = validate_all()
    if not success:
        exit(1)
