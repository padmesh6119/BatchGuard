"""Review-only source/content checks. Never writes to the active project dataset."""
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path

import openpyxl
import pdfplumber


def csv_rows(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def issue_targets(provenance, issues, links):
    """Connect record, document-wide coverage, and linked revision issues."""
    targets = {}
    for issue in issues:
        matches = set()
        for record in provenance:
            direct = record["source_pdf"] == issue["source_pdf"]
            serial = issue["source_serial"]
            if direct and (serial == "N/A" or str(record["source_serial"]) == serial):
                matches.add(record["logical_record_id"])
            for link in links:
                if (issue["source_pdf"] == link["revision_pdf"]
                        and record["output_workbook"] == link["proposed_original_workbook"]
                        and record["output_sheet"] == link["proposed_original_sheet"]
                        and str(record["source_serial"]) == link["proposed_original_serial"]):
                    matches.add(record["logical_record_id"])
        targets[issue["issue_id"]] = matches
    return targets


def annotate_provenance(output_dir):
    output_dir = Path(output_dir)
    path = output_dir / "evidence/record_provenance.jsonl"
    records = [json.loads(line) for line in path.read_text().splitlines()]
    issues = csv_rows(output_dir / "reports/manual_review.csv")
    links = csv_rows(output_dir / "reports/revision_links.csv")
    targets = issue_targets(records, issues, links)
    if any(not values for values in targets.values()):
        raise ValueError("A manual-review issue has no provenance target.")
    by_id = {issue["issue_id"]: issue for issue in issues}
    for record in records:
        record["review_issue_ids"] = sorted(key for key, ids in targets.items()
                                             if record["logical_record_id"] in ids)
        if record["review_issue_ids"]:
            unresolved = any(not by_id[key]["review_status"].startswith("resolved_")
                             for key in record["review_issue_ids"])
            record["parse_status"] = "review_required" if unresolved else "resolved_source_repair"
    path.write_text("".join(json.dumps(record, ensure_ascii=False) + "\n" for record in records))
    return records


def read_source_records(pdf_dir):
    """Re-extract original PDFs without generating or replacing any workbook."""
    from extract_all import extract_standard_table, extract_spurious_table, extract_revision_records
    sources = {}
    for path in sorted(Path(pdf_dir).glob("*.pdf")):
        if "revised" in path.stem:
            continue
        category = path.stem[8:]
        records = (extract_spurious_table(path, path.stem[:7]) if category == "spurious"
                   else extract_standard_table(path, category, path.stem[:7]))
        for record in records:
            key = (path.name, record["serial"])
            if key in sources:
                raise ValueError(f"Duplicate extracted source identity: {key}")
            sources[key] = record
    return sources, extract_revision_records()


def expected_rows(source):
    month = source["source_pdf"][:7]
    if "spurious" not in source["source_pdf"]:
        return [[source["serial"], source["product"], source["batch"], source["mfg_date"],
                 source["exp_date"], source["manufacturer"], source["nsq_result"],
                 "CDSCO Labs" if "central" in source["source_pdf"] else "State Lab",
                 source["lab"], month]]
    # This mirrors the legacy workbook contract. Raw fields stay separate in provenance.
    remark = source["remarks"] or source["reason_for_failure"]
    if not remark:
        raise ValueError("Missing source remark requires explicit review; cannot invent a status.")
    return [[source["serial"], source["product"], source["batch"],
             "Manufacturing Date: " + (source["mfg_date"] or "N/A"), None,
             "CDSCO Labs", source["drawn_by"] or "N/A", month, "NSQ Remark: " + remark],
            [None, None, None, "Expiry Date: " + (source["exp_date"] or "N/A"),
             "Manufactured By: " + (source["manufacturer"] or "Under Investigation"),
             None, None, None, "Firm Reply: " + (source["firm_reply"] or "N/A")]]


def check_integrity(output_dir, pdf_dir, source_data=None):
    """Return concrete errors; passing does not prove the extractor has no remaining bugs."""
    output_dir, pdf_dir = Path(output_dir), Path(pdf_dir)
    sources, revisions = source_data if source_data is not None else read_source_records(pdf_dir)
    errors = []
    hashes = {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in pdf_dir.glob("*.pdf")}
    page_counts = {}
    for path in pdf_dir.glob("*.pdf"):
        with pdfplumber.open(path) as pdf:
            page_counts[path.name] = len(pdf.pages)
    manifest = csv_rows(output_dir / "reports/source_manifest.csv")
    if Counter(row["source_pdf"] for row in manifest) != Counter({key: 1 for key in hashes}):
        errors.append("Manifest does not uniquely cover every source PDF.")
    for row in manifest:
        if row["sha256"] != hashes.get(row["source_pdf"]):
            errors.append("Source manifest hash mismatch: " + row["source_pdf"])
        if str(page_counts.get(row["source_pdf"])) != row["page_count"]:
            errors.append("Source manifest page count mismatch: " + row["source_pdf"])

    provenance = [json.loads(line) for line in (output_dir / "evidence/record_provenance.jsonl").read_text().splitlines()]
    row_key = lambda row: (row["output_workbook"], row["output_sheet"], row["output_excel_row"])
    mapped = {row_key(row): row for row in provenance}
    if len(mapped) != len(provenance):
        errors.append("Duplicate provenance row mapping.")
    actual_keys, encountered = set(), set()
    for path in sorted((output_dir / "workbooks").glob("*.xlsx")):
        workbook = openpyxl.load_workbook(path, data_only=False)
        for sheet in workbook:
            if sheet.merged_cells.ranges:
                errors.append(f"Merged cells: {path.name}/{sheet.title}")
            for cells in sheet.iter_rows(min_row=2):
                key = (path.name, sheet.title, cells[0].row)
                actual_keys.add(key)
                record = mapped.get(key)
                if record is None:
                    errors.append(f"Missing provenance: {key}")
                    continue
                source_key = (record["source_pdf"], record["source_serial"])
                source = sources.get(source_key)
                if source is None:
                    errors.append(f"Missing source record: {key}")
                    continue
                encountered.add(source_key)
                if record["source_pdf"] != path.name[:7] + "_" + sheet.title + ".pdf":
                    errors.append(f"Wrong source/category mapping: {key}")
                if record["source_pdf_sha256"] != hashes.get(record["source_pdf"]):
                    errors.append(f"Provenance source hash mismatch: {key}")
                if record["source_pages"] != source["source_pages"]:
                    errors.append(f"Source page mismatch: {key}")
                if json.loads(record["relevant_raw_text"]) != json.loads(json.dumps(source["raw_rows"])):
                    errors.append(f"Raw PDF table evidence mismatch: {key}")
                expected = expected_rows(source)
                offset = 1 if sheet.title == "spurious" and cells[0].value is None else 0
                normalized = lambda values: [None if value == "" else value for value in values]
                if normalized([cell.value for cell in cells]) != normalized(expected[offset]):
                    errors.append(f"Workbook/source content mismatch: {key}")
                if any(cell.data_type == "f" for cell in cells):
                    errors.append(f"Formula in extracted data: {key}")
                if cells[2].value is not None and cells[2].data_type != "s":
                    errors.append(f"Batch is not text: {key}")
                fields = record["original_extracted_fields"]
                field_map = {"product": "source_product" if source.get("source_repair") else "product",
                             "batch_no": "batch", "mfg_date": "mfg_date", "exp_date": "exp_date",
                             "manufacturer": "manufacturer"}
                field_map.update({name: name for name in ["reason_for_failure", "drawn_by", "firm_reply", "remarks"]}
                                 if sheet.title == "spurious" else {"nsq_result": "nsq_result", "reporting_lab": "lab"})
                if any(fields.get(field) != source.get(source_field) for field, source_field in field_map.items()):
                    errors.append(f"Original fields/source mismatch: {key}")
                if source.get("source_repair") and source["source_repair"] not in record["transformations_performed"]:
                    errors.append(f"Missing source repair annotation: {key}")
        workbook.close()
    if actual_keys != set(mapped):
        errors.append("Provenance row identities do not match workbook row identities.")
    if encountered != set(sources):
        errors.append("Extracted source records and staged logical records differ.")

    issues = csv_rows(output_dir / "reports/manual_review.csv")
    links = csv_rows(output_dir / "reports/revision_links.csv")
    targets = issue_targets(provenance, issues, links)
    if len(targets) != len(issues) or any(not values for values in targets.values()):
        errors.append("Duplicate or unlinked manual-review issue.")
    by_id = {issue["issue_id"]: issue for issue in issues}
    for record in provenance:
        expected_ids = sorted(key for key, ids in targets.items() if record["logical_record_id"] in ids)
        if record["review_issue_ids"] != expected_ids:
            errors.append(f"Missing/incorrect review linkage: {row_key(record)}")
        if expected_ids:
            unresolved = any(not by_id[key]["review_status"].startswith("resolved_") for key in expected_ids)
            expected_status = "review_required" if unresolved else "resolved_source_repair"
            if record["parse_status"] != expected_status:
                errors.append(f"Review status mismatch: {row_key(record)}")
    for issue_id, serial in [("REV-004", 25), ("REV-005", 26)]:
        issue = by_id.get(issue_id, {})
        source = sources[("2025-06_central.pdf", serial)]
        text = " ".join(issue.get("raw_evidence", "").split())
        if " ".join(source["product"].split()) not in text or "Ibuprofen" in text:
            errors.append(f"Incorrect manual-review product description: {issue_id}")
    if "Vitamin B Complex Tablets N.F.I" not in by_id.get("REV-003", {}).get("proposed_action", ""):
        errors.append("Incorrect product repair description: REV-003")

    summary = csv_rows(output_dir / "reports/extraction_summary.csv")
    if Counter(row["source_pdf"] for row in summary) != Counter({key: 1 for key in hashes}):
        errors.append("Summary does not uniquely cover every source PDF.")
    for row in summary:
        name = row["source_pdf"]
        revision = "revised" in name
        expected_count = sum(record["revision_pdf"] == name for record in revisions) if revision else sum(key[0] == name for key in sources)
        if str(expected_count) != row["source_logical_entries"]:
            errors.append("Summary source count mismatch: " + name)
        if row["workbook_logical_records"] != str(0 if revision else expected_count):
            errors.append("Summary workbook count mismatch: " + name)
        pending = {record["logical_record_id"] for record in provenance
                   if record["source_pdf"] == name and record["parse_status"] == "review_required"}
        if revision:
            pending = {name} if any(issue["source_pdf"] == name and not issue["review_status"].startswith("resolved_") for issue in issues) else set()
        if row["unresolved_entries"] != str(len(pending)):
            errors.append("Summary unresolved count mismatch: " + name)

    revision_book = openpyxl.load_workbook(output_dir / "evidence/revision_records.xlsx", data_only=False)
    revision_rows = list(revision_book["revision_records"].values)[1:]
    revision_keys = ["revision_pdf", "report_month_context", "report_heading", "report_subtitle", "serial",
                     "product", "batch", "mfg_date", "exp_date", "manufacturer", "reason_for_failure",
                     "drawn_by", "firm_reply", "remarks"]
    expected_revision_rows = [tuple(record[key] or None for key in revision_keys) for record in revisions]
    if revision_rows != expected_revision_rows:
        errors.append("Revision workbook differs from re-extracted revision PDFs.")
    revision_book.close()
    revision_map = {record["revision_pdf"]: record for record in revisions}
    if Counter(link["revision_pdf"] for link in links) != Counter({key: 1 for key in revision_map}):
        errors.append("Revision links do not uniquely cover all revisions.")
    for link in links:
        source_name = link["proposed_original_workbook"][:7] + "_" + link["proposed_original_sheet"] + ".pdf"
        original = sources.get((source_name, int(link["proposed_original_serial"])))
        revision = revision_map.get(link["revision_pdf"])
        compact = lambda text: "".join(text.split())
        if not original or not revision or compact(original["batch"]) != compact(revision["batch"]):
            errors.append("Revision link batch/source mismatch: " + link["revision_pdf"])
        elif compact(link["revision_batch"]) != compact(revision["batch"]):
            errors.append("Revision link description mismatch: " + link["revision_pdf"])
    return errors
