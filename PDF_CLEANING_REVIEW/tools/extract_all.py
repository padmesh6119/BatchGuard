#!/usr/bin/env python3
"""
tools/extract_all.py
Deterministic PDF extraction tool for BatchGuard earlier CDSCO reports (Sep 2024 - Jun 2025).
Extracts 10 central reports, 10 state reports, 8 spurious reports, and 4 revision documents.
Outputs:
  - 10 monthly workbooks under PDF_CLEANING_REVIEW/workbooks/
  - evidence/revision_records.xlsx
  - evidence/record_provenance.jsonl
"""

import json
import re
import hashlib
from pathlib import Path
import pdfplumber
import openpyxl
from openpyxl.cell import Cell

PROJECT_ROOT = Path("/home/vishu/Projects/IBM/BATCHGUARD")
PDF_DIR = PROJECT_ROOT / "PDF"
OUTPUT_DIR = PROJECT_ROOT / "PDF_CLEANING_REVIEW"
WORKBOOKS_DIR = OUTPUT_DIR / "workbooks"
EVIDENCE_DIR = OUTPUT_DIR / "evidence"
REPORTS_DIR = OUTPUT_DIR / "reports"
TOOLS_DIR = OUTPUT_DIR / "tools"

MONTHS = [
    "2024-09", "2024-10", "2024-11", "2024-12",
    "2025-01", "2025-02", "2025-03", "2025-04",
    "2025-05", "2025-06"
]

CENTRAL_HEADERS = [
    "S.No",
    "Name of Product",
    "Batch No",
    "Manufacturing Date",
    "Expiry Date",
    "Manufactured By",
    "NSQ Result",
    "Reporting Source",
    "Reporting by Lab/State",
    "Reporting Month & Year",
]

STATE_HEADERS = [
    "S.No",
    "Name of Product",
    "Batch No",
    "Manufacturing Date",
    "Expiry Date",
    "Manufactured By",
    "NSQ Result",
    "Reporting Source",
    "Reporting by Lab/State",
    "Reporting Month & Year",
]

SPURIOUS_HEADERS = [
    "S.No",
    "Name of Product",
    "Batch No",
    "Manufacturing Dates",
    "Manufacturer Details",
    "Reporting Source",
    "Reporting by Lab/State",
    "Reporting Month & Year",
    "Remarks",
]


def clean_cell_text(text: str) -> str:
    """Normalize extracted cell text while preserving internal wording."""
    if text is None:
        return ""
    text = str(text).replace("\xa0", " ").strip()
    # Normalize multiple whitespaces on a single line, but preserve newlines
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    return "\n".join(l for l in lines if l)


def get_file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_excel_row(sheet, row_idx: int, values: list, force_text_cols: set = None):
    """Write a row of values into an openpyxl sheet with explicit text formatting for batch/serials."""
    for col_idx, val in enumerate(values, 1):
        cell = sheet.cell(row=row_idx, column=col_idx)
        if val is None or val == "":
            cell.value = None
            continue
        if force_text_cols and (col_idx - 1) in force_text_cols:
            cell.number_format = "@"
            cell.data_type = "s"
            cell.value = str(val)
        else:
            if isinstance(val, int):
                cell.value = val
            else:
                cell.number_format = "@"
                cell.data_type = "s"
                cell.value = str(val)


def extract_standard_table(pdf_path: Path, category_name: str, report_month: str):
    """
    Extract central or state report.
    Returns list of logical records and raw metadata.
    """
    records = []
    curr = None
    curr_pages = []
    sha256 = get_file_sha256(pdf_path)
    
    with pdfplumber.open(pdf_path) as pdf:
        for p_idx, page in enumerate(pdf.pages, 1):
            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    if not row or all(c is None or str(c).strip() == "" for c in row):
                        continue
                    if len(row) < 5:
                        # Non-record footnote, title, or boilerplate
                        continue
                    
                    c0 = clean_cell_text(row[0])
                    # Skip repeated table headers
                    if any(k in c0.lower() for k in ["s.no", "s. no", "cdsco", "state", "product/drug"]):
                        continue
                    
                    # Match serial number at start of c0
                    m = re.match(r"^(\d+)\.?(?:\s*(.*))?$", c0, re.DOTALL)
                    if m:
                        if curr:
                            curr["source_pages"] = sorted(set(curr_pages))
                            records.append(curr)
                        
                        serial_num = int(m.group(1))
                        overflow = m.group(2).strip() if m.group(2) else ""
                        curr_pages = [p_idx]
                        
                        prod = clean_cell_text(row[1])
                        source_product = prod
                        if overflow:
                            prod = (overflow + " " + prod).strip()
                        repair = None
                        if pdf_path.name == "2024-10_state.pdf" and serial_num == 13:
                            # Verified against page 5 and independent Poppler text.
                            # Co belongs before mplex, not before Vitamin B.
                            if overflow != "Co" or " ".join(source_product.split()) != "Vitamin B mplex Tablets N.F.I":
                                raise ValueError("Reviewed product-overflow source changed; inspect page 5 again.")
                            prod = "Vitamin B Complex Tablets N.F.I"
                            repair = "verified_source_repair_REV-003"
                            
                        curr = {
                            "serial": serial_num,
                            "product": prod,
                            "source_product": source_product,
                            "source_repair": repair,
                            "batch": clean_cell_text(row[2]),
                            "mfg_date": clean_cell_text(row[3]),
                            "exp_date": clean_cell_text(row[4]),
                            "manufacturer": clean_cell_text(row[5]),
                            "nsq_result": clean_cell_text(row[6]),
                            "lab": clean_cell_text(row[7]) if len(row) > 7 else "",
                            "source_pages": curr_pages,
                            "source_pdf": pdf_path.name,
                            "source_pdf_sha256": sha256,
                            "overflow_text": overflow,
                            "raw_rows": [(p_idx, [clean_cell_text(c) for c in row])],
                        }
                    elif curr:
                        # Continuation row from wrapped cell
                        curr_pages.append(p_idx)
                        curr["raw_rows"].append((p_idx, [clean_cell_text(c) for c in row]))
                        
                        field_map = [
                            (1, "product"),
                            (2, "batch"),
                            (3, "mfg_date"),
                            (4, "exp_date"),
                            (5, "manufacturer"),
                            (6, "nsq_result"),
                        ]
                        for idx, key in field_map:
                            if len(row) > idx and row[idx]:
                                val = clean_cell_text(row[idx])
                                if val:
                                    curr[key] = (curr[key] + " " + val).strip() if curr[key] else val
                        if len(row) > 7 and row[7]:
                            val = clean_cell_text(row[7])
                            if val:
                                curr["lab"] = (curr["lab"] + " " + val).strip() if curr["lab"] else val

        if curr:
            curr["source_pages"] = sorted(set(curr_pages))
            records.append(curr)
            
    return records


def extract_spurious_table(pdf_path: Path, report_month: str):
    """
    Extract spurious report into structured logical records.
    """
    records = []
    sha256 = get_file_sha256(pdf_path)
    
    with pdfplumber.open(pdf_path) as pdf:
        all_raw_rows = []
        for p_idx, page in enumerate(pdf.pages, 1):
            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    if not row or all(c is None or str(c).strip() == "" for c in row):
                        continue
                    all_raw_rows.append((p_idx, row))
                    
    curr = None
    curr_pages = []
    
    # Check if this PDF has 9 or 10 standard columns or 12 (like 2025-06)
    for p_idx, row in all_raw_rows:
        if len(row) < 5:
            continue
        c0 = clean_cell_text(row[0])
        if any(k in c0.lower() for k in ["s.no", "s. no", "drugs/medical"]):
            continue
            
        m = re.match(r"^(\d+)\.?(?:\s*(.*))?$", c0, re.DOTALL)
        if m:
            if curr:
                curr["source_pages"] = sorted(set(curr_pages))
                records.append(curr)
                
            serial_num = int(m.group(1))
            overflow = m.group(2).strip() if m.group(2) else ""
            curr_pages = [p_idx]
            
            # Normalize column indexes based on row length
            if len(row) >= 12:
                # 2025-06 spurious page 1 split product across cols 1, 2, 3
                prod = " ".join(clean_cell_text(c) for c in row[1:4] if clean_cell_text(c)).strip()
                batch = clean_cell_text(row[4])
                mfg = clean_cell_text(row[5])
                exp = clean_cell_text(row[6])
                mfr = clean_cell_text(row[7])
                reason = clean_cell_text(row[8])
                drawn = clean_cell_text(row[9])
                firm = clean_cell_text(row[10])
                remarks = clean_cell_text(row[11]) if len(row) > 11 else ""
            elif len(row) == 10:
                prod = clean_cell_text(row[1])
                batch = clean_cell_text(row[2])
                mfg = clean_cell_text(row[3])
                exp = clean_cell_text(row[4])
                mfr = clean_cell_text(row[5])
                reason = clean_cell_text(row[6])
                drawn = clean_cell_text(row[7])
                firm = clean_cell_text(row[8])
                remarks = clean_cell_text(row[9])
            else: # 9 columns (2024-09)
                prod = clean_cell_text(row[1])
                batch = clean_cell_text(row[2])
                mfg = clean_cell_text(row[3])
                exp = clean_cell_text(row[4])
                mfr = clean_cell_text(row[5])
                reason = clean_cell_text(row[6])
                drawn = clean_cell_text(row[7])
                firm = clean_cell_text(row[8])
                remarks = ""
                
            if overflow:
                prod = (overflow + " " + prod).strip()
                
            curr = {
                "serial": serial_num,
                "product": prod,
                "batch": batch,
                "mfg_date": mfg,
                "exp_date": exp,
                "manufacturer": mfr,
                "reason_for_failure": reason,
                "drawn_by": drawn,
                "firm_reply": firm,
                "remarks": remarks,
                "source_pages": curr_pages,
                "source_pdf": pdf_path.name,
                "source_pdf_sha256": sha256,
                "raw_rows": [(p_idx, [clean_cell_text(c) for c in row])],
            }
        elif curr:
            # Continuation row
            curr_pages.append(p_idx)
            curr["raw_rows"].append((p_idx, [clean_cell_text(c) for c in row]))
            
            if len(row) >= 12:
                prod_part = " ".join(clean_cell_text(c) for c in row[1:4] if clean_cell_text(c)).strip()
                if prod_part:
                    curr["product"] = (curr["product"] + " " + prod_part).strip()
                for idx, key in [(4, "batch"), (5, "mfg_date"), (6, "exp_date"), (7, "manufacturer"),
                                 (8, "reason_for_failure"), (9, "drawn_by"), (10, "firm_reply"), (11, "remarks")]:
                    if len(row) > idx and row[idx]:
                        val = clean_cell_text(row[idx])
                        if val:
                            curr[key] = (curr[key] + " " + val).strip() if curr[key] else val
            elif len(row) == 10:
                for idx, key in [(1, "product"), (2, "batch"), (3, "mfg_date"), (4, "exp_date"), (5, "manufacturer"),
                                 (6, "reason_for_failure"), (7, "drawn_by"), (8, "firm_reply"), (9, "remarks")]:
                    if len(row) > idx and row[idx]:
                        val = clean_cell_text(row[idx])
                        if val:
                            curr[key] = (curr[key] + " " + val).strip() if curr[key] else val
            else:
                for idx, key in [(1, "product"), (2, "batch"), (3, "mfg_date"), (4, "exp_date"), (5, "manufacturer"),
                                 (6, "reason_for_failure"), (7, "drawn_by"), (8, "firm_reply")]:
                    if len(row) > idx and row[idx]:
                        val = clean_cell_text(row[idx])
                        if val:
                            curr[key] = (curr[key] + " " + val).strip() if curr[key] else val

    if curr:
        curr["source_pages"] = sorted(set(curr_pages))
        records.append(curr)
        
    return records


def extract_revision_records():
    """Extract records from all 4 revision PDFs."""
    revision_files = [
        ("2024-11_revised.pdf", "2024-11"),
        ("2025-01_revised.pdf", "2025-01"),
        ("2025-01_revised-1.pdf", "2025-01"),
        ("2025-05_revised.pdf", "2025-05"),
    ]
    
    extracted = []
    for fname, month in revision_files:
        path = PDF_DIR / fname
        sha256 = get_file_sha256(path)
        with pdfplumber.open(path) as pdf:
            p0 = pdf.pages[0]
            text = p0.extract_text()
            lines = text.split("\n")
            heading = lines[0] if lines else ""
            subtitle = lines[1] if len(lines) > 1 else ""
            table = p0.extract_table()
            
            for row in table[1:]:
                if not row or all(c is None or str(c).strip() == "" for c in row):
                    continue
                c0 = clean_cell_text(row[0])
                m = re.match(r"^(\d+)\.?", c0)
                serial = int(m.group(1)) if m else c0
                
                record = {
                    "revision_pdf": fname,
                    "revision_pdf_sha256": sha256,
                    "report_month_context": month,
                    "report_heading": heading,
                    "report_subtitle": subtitle,
                    "serial": serial,
                    "product": clean_cell_text(row[1]),
                    "batch": clean_cell_text(row[2]),
                    "mfg_date": clean_cell_text(row[3]),
                    "exp_date": clean_cell_text(row[4]),
                    "manufacturer": clean_cell_text(row[5]),
                    "reason_for_failure": clean_cell_text(row[6]),
                    "drawn_by": clean_cell_text(row[7]),
                    "firm_reply": clean_cell_text(row[8]),
                    "remarks": clean_cell_text(row[9]) if len(row) > 9 else "",
                }
                extracted.append(record)
    return extracted


def build_workbooks():
    """Main builder for all workbooks and evidence outputs."""
    WORKBOOKS_DIR.mkdir(parents=True, exist_ok=True)
    EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    
    provenance_file = EVIDENCE_DIR / "record_provenance.jsonl"
    provenance_entries = []
    
    workbook_stats = {}
    
    for month in MONTHS:
        wb_filename = f"{month}_all.xlsx"
        wb_path = WORKBOOKS_DIR / wb_filename
        wb = openpyxl.Workbook()
        
        # 1. Central Sheet
        ws_central = wb.active
        ws_central.title = "central"
        write_excel_row(ws_central, 1, CENTRAL_HEADERS)
        
        central_pdf = PDF_DIR / f"{month}_central.pdf"
        central_records = extract_standard_table(central_pdf, "central", month)
        
        for idx, rec in enumerate(central_records, 2):
            excel_row = [
                rec["serial"],
                rec["product"],
                rec["batch"],
                rec["mfg_date"],
                rec["exp_date"],
                rec["manufacturer"],
                rec["nsq_result"],
                "CDSCO Labs",
                rec["lab"],
                month,
            ]
            # force text formatting on Batch No (col 2, 0-indexed)
            write_excel_row(ws_central, idx, excel_row, force_text_cols={2})
            
            # Provenance
            provenance_entries.append({
                "output_workbook": wb_filename,
                "output_sheet": "central",
                "output_excel_row": idx,
                "logical_record_id": f"{month}:central:{rec['serial']:04d}",
                "source_pdf": rec["source_pdf"],
                "source_pdf_sha256": rec["source_pdf_sha256"],
                "source_pages": rec["source_pages"],
                "source_serial": rec["serial"],
                "original_extracted_fields": {
                    "product": rec["source_product"] if rec.get("source_repair") else rec["product"],
                    "batch_no": rec["batch"],
                    "mfg_date": rec["mfg_date"],
                    "exp_date": rec["exp_date"],
                    "manufacturer": rec["manufacturer"],
                    "nsq_result": rec["nsq_result"],
                    "reporting_lab": rec["lab"],
                },
                "relevant_raw_text": json.dumps(rec["raw_rows"], ensure_ascii=False),
                "transformations_performed": [
                    "joined_wrapped_cells",
                    "trimmed_whitespace",
                    "formatted_batch_as_text",
                    f"assigned_reporting_source_CDSCO_Labs",
                    f"assigned_reporting_month_{month}",
                ],
                "parse_status": "clean" if not rec.get("overflow_text") else "repaired_overflow",
                "review_issue_ids": [],
            })
            
        # 2. State Sheet
        ws_state = wb.create_sheet(title="state")
        write_excel_row(ws_state, 1, STATE_HEADERS)
        
        state_pdf = PDF_DIR / f"{month}_state.pdf"
        state_records = extract_standard_table(state_pdf, "state", month)
        
        for idx, rec in enumerate(state_records, 2):
            excel_row = [
                rec["serial"],
                rec["product"],
                rec["batch"],
                rec["mfg_date"],
                rec["exp_date"],
                rec["manufacturer"],
                rec["nsq_result"],
                "State Lab",
                rec["lab"],
                month,
            ]
            write_excel_row(ws_state, idx, excel_row, force_text_cols={2})
            
            provenance_entries.append({
                "output_workbook": wb_filename,
                "output_sheet": "state",
                "output_excel_row": idx,
                "logical_record_id": f"{month}:state:{rec['serial']:04d}",
                "source_pdf": rec["source_pdf"],
                "source_pdf_sha256": rec["source_pdf_sha256"],
                "source_pages": rec["source_pages"],
                "source_serial": rec["serial"],
                "original_extracted_fields": {
                    "product": rec["source_product"] if rec.get("source_repair") else rec["product"],
                    "batch_no": rec["batch"],
                    "mfg_date": rec["mfg_date"],
                    "exp_date": rec["exp_date"],
                    "manufacturer": rec["manufacturer"],
                    "nsq_result": rec["nsq_result"],
                    "reporting_lab": rec["lab"],
                },
                "relevant_raw_text": json.dumps(rec["raw_rows"], ensure_ascii=False),
                "transformations_performed": [
                    "joined_wrapped_cells",
                    "trimmed_whitespace",
                    "formatted_batch_as_text",
                    "assigned_reporting_source_State_Lab",
                    f"assigned_reporting_month_{month}",
                ],
                "parse_status": "clean" if not rec.get("overflow_text") else "repaired_overflow",
                "review_issue_ids": [],
            })
            
        # 3. Spurious Sheet (ONLY if standalone spurious source exists)
        spurious_pdf = PDF_DIR / f"{month}_spurious.pdf"
        spurious_records = []
        if spurious_pdf.exists():
            ws_spurious = wb.create_sheet(title="spurious")
            write_excel_row(ws_spurious, 1, SPURIOUS_HEADERS)
            spurious_records = extract_spurious_table(spurious_pdf, month)
            
            excel_curr_row = 2
            for rec in spurious_records:
                logical_rec_id = f"{month}:spurious:{rec['serial']:04d}"
                
                # Format 2-row pattern
                mfg_val = rec["mfg_date"] or "N/A"
                exp_val = rec["exp_date"] or "N/A"
                mfr_val = rec["manufacturer"] or "Under Investigation"
                
                # Remarks formatting
                if rec["remarks"]:
                    nsq_remark_text = f"NSQ Remark: {rec['remarks']}"
                elif rec["reason_for_failure"]:
                    nsq_remark_text = f"NSQ Remark: {rec['reason_for_failure']}"
                else:
                    nsq_remark_text = "NSQ Remark: Purported to be spurious, subject to outcome of investigation."
                    
                firm_reply_text = f"Firm Reply: {rec['firm_reply']}" if rec["firm_reply"] else "Firm Reply: N/A"
                
                # Row 1 (Primary)
                row_1 = [
                    rec["serial"],
                    rec["product"],
                    rec["batch"],
                    f"Manufacturing Date: {mfg_val}",
                    None, # Manufacturer Name ONLY when explicit name given separately
                    "CDSCO Labs",
                    rec["drawn_by"] or "N/A",
                    month,
                    nsq_remark_text,
                ]
                write_excel_row(ws_spurious, excel_curr_row, row_1, force_text_cols={2})
                
                # Provenance for Row 1
                provenance_entries.append({
                    "output_workbook": wb_filename,
                    "output_sheet": "spurious",
                    "output_excel_row": excel_curr_row,
                    "logical_record_id": logical_rec_id,
                    "source_pdf": rec["source_pdf"],
                    "source_pdf_sha256": rec["source_pdf_sha256"],
                    "source_pages": rec["source_pages"],
                    "source_serial": rec["serial"],
                    "original_extracted_fields": {
                        "product": rec["product"],
                        "batch_no": rec["batch"],
                        "mfg_date": rec["mfg_date"],
                        "exp_date": rec["exp_date"],
                        "manufacturer": rec["manufacturer"],
                        "reason_for_failure": rec["reason_for_failure"],
                        "drawn_by": rec["drawn_by"],
                        "firm_reply": rec["firm_reply"],
                        "remarks": rec["remarks"],
                    },
                    "relevant_raw_text": json.dumps(rec["raw_rows"], ensure_ascii=False),
                    "transformations_performed": [
                        "formatted_two_row_spurious_pattern_row_1",
                        "formatted_batch_as_text",
                        "prefixed_manufacturing_date",
                        "prefixed_nsq_remark",
                    ],
                    "parse_status": "clean",
                    "review_issue_ids": [],
                })
                excel_curr_row += 1
                
                # Row 2 (Continuation)
                row_2 = [
                    None,
                    None,
                    None,
                    f"Expiry Date: {exp_val}",
                    f"Manufactured By: {mfr_val}",
                    None,
                    None,
                    None,
                    firm_reply_text,
                ]
                write_excel_row(ws_spurious, excel_curr_row, row_2)
                
                # Provenance for Row 2
                provenance_entries.append({
                    "output_workbook": wb_filename,
                    "output_sheet": "spurious",
                    "output_excel_row": excel_curr_row,
                    "logical_record_id": logical_rec_id,
                    "source_pdf": rec["source_pdf"],
                    "source_pdf_sha256": rec["source_pdf_sha256"],
                    "source_pages": rec["source_pages"],
                    "source_serial": rec["serial"],
                    "original_extracted_fields": {
                        "product": rec["product"],
                        "batch_no": rec["batch"],
                        "mfg_date": rec["mfg_date"],
                        "exp_date": rec["exp_date"],
                        "manufacturer": rec["manufacturer"],
                        "reason_for_failure": rec["reason_for_failure"],
                        "drawn_by": rec["drawn_by"],
                        "firm_reply": rec["firm_reply"],
                        "remarks": rec["remarks"],
                    },
                    "relevant_raw_text": json.dumps(rec["raw_rows"], ensure_ascii=False),
                    "transformations_performed": [
                        "formatted_two_row_spurious_pattern_row_2_continuation",
                        "prefixed_expiry_date",
                        "prefixed_manufactured_by",
                        "prefixed_firm_reply",
                    ],
                    "parse_status": "clean",
                    "review_issue_ids": [],
                })
                excel_curr_row += 1
                
        wb.save(wb_path)
        workbook_stats[month] = {
            "central": len(central_records),
            "state": len(state_records),
            "spurious": len(spurious_records),
        }
        print(f"Generated {wb_filename:<16} | central: {len(central_records):>3} | state: {len(state_records):>3} | spurious: {len(spurious_records):>2}")

    # Write provenance JSONL
    for entry in provenance_entries:
        if entry["source_pdf"] == "2024-10_state.pdf" and entry["source_serial"] == 13:
            entry["transformations_performed"].append("verified_source_repair_REV-003")
    with open(provenance_file, "w", encoding="utf-8") as f:
        for entry in provenance_entries:
            f.write(json.dumps(entry, ensure_ascii=False) + "\n")
    print(f"Saved provenance records: {len(provenance_entries)} entries in {provenance_file}")
    
    # Generate revision_records.xlsx
    revision_records = extract_revision_records()
    wb_rev = openpyxl.Workbook()
    ws_rev = wb_rev.active
    ws_rev.title = "revision_records"
    
    rev_headers = [
        "revision_pdf", "report_month_context", "report_heading", "report_subtitle",
        "S.No", "Name of Product", "Batch No", "Date of Manufacture", "Date of Expiry",
        "Manufactured By", "Reason for failure", "Drawn By", "Firm Reply", "Remarks"
    ]
    write_excel_row(ws_rev, 1, rev_headers)
    for r_idx, r in enumerate(revision_records, 2):
        row_vals = [
            r["revision_pdf"], r["report_month_context"], r["report_heading"], r["report_subtitle"],
            r["serial"], r["product"], r["batch"], r["mfg_date"], r["exp_date"],
            r["manufacturer"], r["reason_for_failure"], r["drawn_by"], r["firm_reply"], r["remarks"]
        ]
        write_excel_row(ws_rev, r_idx, row_vals, force_text_cols={6})
    rev_path = EVIDENCE_DIR / "revision_records.xlsx"
    wb_rev.save(rev_path)
    print(f"Saved {len(revision_records)} revision records in {rev_path}")


if __name__ == "__main__":
    build_workbooks()
